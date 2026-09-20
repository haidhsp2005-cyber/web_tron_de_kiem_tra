import os
import re
from copy import deepcopy
import docx
from docx.oxml.ns import qn
from docx.oxml import parse_xml
from .formula_helper import (
    is_run_element_red,
    is_math_element_red,
    extract_element_text_with_formatting,
    serialize_oxml_elements,
    deserialize_oxml_elements,
    strip_elements_prefix,
    strip_elements_suffix,
    heal_omath_element
)

PART1_REGEX = re.compile(r"^\s*PHẦN\s+(I|1)[\.\:\s\-]", re.IGNORECASE)
PART2_REGEX = re.compile(r"^\s*PHẦN\s+(II|2)[\.\:\s\-]", re.IGNORECASE)
PART3_REGEX = re.compile(r"^\s*PHẦN\s+(III|3)[\.\:\s\-]", re.IGNORECASE)
PART4_REGEX = re.compile(r"^\s*PHẦN\s+(IV|4)[\.\:\s\-]", re.IGNORECASE)

QUESTION_REGEX = re.compile(r"^\s*(Câu|Bài)\s+(\d+)[\.\:\-\s]", re.IGNORECASE)
CHOICE_SPLIT_REGEX = re.compile(r"(?=(?:^|\s{2,}|\t)([A-D])[\.\:\)])")

class DocxParser:
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.doc = docx.Document(file_path)
        self.metadata = {
            "title": "ĐỀ KIỂM TRA",
            "school": "THPT Long Cang",
            "subject": "Toán",
            "time": "90 phút",
            "year": "2026 - 2027"
        }
        self.part1_questions = []
        self.part2_questions = []
        self.part3_questions = []
        self.part4_questions = []

    def parse(self) -> dict:
        """Parses the DOCX file and extracts sections and questions."""
        paragraphs_data = self._extract_paragraphs_with_runs()
        
        current_part = 1
        i = 0
        n = len(paragraphs_data)
        
        while i < n:
            p_data = paragraphs_data[i]
            text = p_data["raw_text"].strip()
            
            # Skip instructions box or headers
            if "HƯỚNG DẪN ĐỊNH DẠNG" in text.upper():
                i += 1
                continue
                
            # Detect Answer Key Section at document end (Stops exam question parsing)
            if (len(self.part1_questions) > 0 or len(self.part2_questions) > 0) and re.match(r"^\s*(BẢNG\s+ĐÁP\s+ÁN|HƯỚNG\s+DẪN\s+CHẤM\s+ĐỀ|ĐÁP\s+ÁN\s+VÀ\s+THANG\s+ĐIỂM|BẢNG\s+TRẢ\s+LỜI)", text, re.IGNORECASE):
                self._parse_trailing_answers(paragraphs_data, i)
                break

            # Detect Sections
            if PART1_REGEX.search(text):
                current_part = 1
                i += 1
                continue
            elif PART2_REGEX.search(text):
                current_part = 2
                i += 1
                continue
            elif PART3_REGEX.search(text):
                current_part = 3
                i += 1
                continue
            elif PART4_REGEX.search(text):
                current_part = 4
                i += 1
                continue
                
            # Detect Questions
            q_match = QUESTION_REGEX.search(text)
            if q_match:
                if current_part == 1:
                    i = self._parse_part1_question(paragraphs_data, i)
                elif current_part == 2:
                    i = self._parse_part2_question(paragraphs_data, i)
                elif current_part == 3:
                    i = self._parse_part3_question(paragraphs_data, i)
                elif current_part == 4:
                    i = self._parse_part4_question(paragraphs_data, i)
                else:
                    i += 1
            else:
                # Could be header metadata
                self._check_metadata(text)
                i += 1

        # Check if any answer tables in doc.tables were not processed
        self._check_trailing_tables()

        return {
            "metadata": self.metadata,
            "part1": self.part1_questions,
            "part2": self.part2_questions,
            "part3": self.part3_questions,
            "part4": self.part4_questions,
            "stats": {
                "total_questions": len(self.part1_questions) + len(self.part2_questions) + len(self.part3_questions) + len(self.part4_questions),
                "part1_count": len(self.part1_questions),
                "part2_count": len(self.part2_questions),
                "part3_count": len(self.part3_questions),
                "part4_count": len(self.part4_questions)
            }
        }

    def _extract_paragraphs_with_runs(self) -> list[dict]:
        """Extracts text, formatted HTML, red flags and oxml elements from document paragraphs & tables in exact document order."""
        items = []
        body = self.doc.element.body
        for child in body.iterchildren():
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "p":
                p = docx.text.paragraph.Paragraph(child, self.doc)
                item = self._process_paragraph(p)
                if item["raw_text"].strip():
                    items.append(item)
            elif tag == "tbl":
                tbl = docx.table.Table(child, self.doc)
                for row in tbl.rows:
                    for cell in row.cells:
                        for cp in cell.paragraphs:
                            item = self._process_paragraph(cp)
                            if item["raw_text"].strip():
                                items.append(item)
        return items

    def _process_paragraph(self, p) -> dict:
        formatted_parts = []
        raw_parts = []
        oxml_elements = []
        has_red = False

        for child in p._p:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag in ["r", "oMath", "oMathPara"]:
                heal_omath_element(child)
                text, is_red = extract_element_text_with_formatting(child)
                if text:
                    formatted_parts.append(text)
                    raw_text_part = re.sub(r"<br\s*/?>", "\n", text)
                    raw_parts.append(re.sub(r"<[^>]+>", "", raw_text_part))
                if is_red:
                    has_red = True
                oxml_elements.append(child)
            elif tag in ["br", "cr"]:
                formatted_parts.append("<br>")
                raw_parts.append("\n")
                oxml_elements.append(child)

        return {
            "raw_text": "".join(raw_parts),
            "formatted_text": "".join(formatted_parts),
            "has_red": has_red,
            "oxml_elements": oxml_elements,
            "xml_strings": serialize_oxml_elements(oxml_elements)
        }

    def _check_metadata(self, text: str):
        if "TRƯỜNG" in text.upper():
            m = re.search(r"(TRƯỜNG\s+THPT\s+[A-ZÀ-Ỹa-zà-ỹ0-9\s]+)", text, re.IGNORECASE)
            if m:
                clean_sch = re.split(r"NĂM\s*HỌC|KỲ\s*THI|[\n\r,]", m.group(1), flags=re.IGNORECASE)[0].strip()
                if clean_sch:
                    self.metadata["school"] = clean_sch
            else:
                lines = [l.strip() for l in text.split("\n") if "TRƯỜNG" in l.upper()]
                if lines:
                    clean_sch = re.split(r"NĂM\s*HỌC|KỲ\s*THI|[\n\r,]", lines[0], flags=re.IGNORECASE)[0].strip()
                    if clean_sch:
                        self.metadata["school"] = clean_sch

        if "NĂM HỌC" in text.upper() or re.search(r"20\d\d\s*[\-\–]\s*20\d\d", text):
            m = re.search(r"(20\d\d\s*[\-\–]\s*20\d\d)", text)
            if m:
                self.metadata["year"] = m.group(1).replace("–", "-").strip()

        if "MÔN:" in text.upper() or "MÔN " in text.upper():
            m = re.search(r"MÔN[\:\s]+([^\n\r,\.\-]+)", text, re.IGNORECASE)
            if m:
                self.metadata["subject"] = m.group(1).strip()
        if "THỜI GIAN" in text.upper():
            m = re.search(r"(\d+)\s*phút", text, re.IGNORECASE)
            if m:
                self.metadata["time"] = f"{m.group(1)} phút"


    # ================= PART 1 PARSING =================
    def _parse_part1_question(self, paragraphs_data, start_idx) -> int:
        p_data = paragraphs_data[start_idx]
        q_text = p_data["formatted_text"]
        raw_text = p_data["raw_text"]
        
        # Clean question prefix (e.g. "Câu 1: ")
        q_match = QUESTION_REGEX.search(raw_text)
        q_num = int(q_match.group(2)) if q_match else len(self.part1_questions) + 1
        
        # Question content
        clean_q_text = re.sub(r"^\s*(Câu|Bài)\s+\d+[\.\:\-\s]+", "", q_text).strip()
        clean_q_xmls = strip_elements_prefix(p_data["xml_strings"], r"^\s*(Câu|Bài)\s+\d+[\.\:\-\s]*")
        
        choices = {"A": "", "B": "", "C": "", "D": ""}
        correct_answer = "A" # Default if not found
        choice_xmls = {"A": [], "B": [], "C": [], "D": []}
        found_red = False

        # Look forward for choices
        idx = start_idx + 1
        # First check if choices were in the same paragraph
        inline_choices = self._extract_choices_from_p(p_data)
        if len(inline_choices) >= 2:
            choices.update(inline_choices)
            for k, c_info in inline_choices.items():
                if c_info.get("is_red"):
                    correct_answer = k
                    found_red = True
            self.part1_questions.append({
                "id": len(self.part1_questions) + 1,
                "original_num": q_num,
                "question": clean_q_text,
                "choices": {k: v.get("text", "") for k, v in choices.items()},
                "choice_xmls": {k: v.get("xml_strings", []) for k, v in choices.items()},
                "correct": correct_answer,
                "has_red": found_red,
                "xml_strings": clean_q_xmls
            })
            return idx

        seen_choice_keys = set()
        active_choice = None

        # Otherwise look at succeeding paragraphs
        while idx < len(paragraphs_data):
            next_p = paragraphs_data[idx]
            next_text = next_p["raw_text"].strip()
            
            # If hit another question, section, or answer key / end marker, stop
            if (QUESTION_REGEX.search(next_text) or PART1_REGEX.search(next_text) or 
                PART2_REGEX.search(next_text) or PART3_REGEX.search(next_text) or 
                PART4_REGEX.search(next_text) or 
                re.search(r"(BẢNG\s+ĐÁP\s+ÁN|HƯỚNG\s+DẪN\s+CHẤM|\-\-\-+\s*HẾT)", next_text, re.IGNORECASE)):
                break
                
            p_choices = self._extract_choices_from_p(next_p)
            if p_choices:
                for k, c_info in p_choices.items():
                    seen_choice_keys.add(k)
                    choices[k] = c_info.get("text", "")
                    choice_xmls[k] = c_info.get("xml_strings", [])
                    if c_info.get("is_red"):
                        correct_answer = k
                        found_red = True
                active_choice = sorted(p_choices.keys())[-1]
                idx += 1
            else:
                if not seen_choice_keys:
                    clean_q_text += "<br>" + next_p["formatted_text"]
                    clean_q_xmls.extend(next_p["xml_strings"])
                    idx += 1
                elif active_choice:
                    choices[active_choice] = (choices[active_choice] + " " + next_p["formatted_text"]).strip()
                    choice_xmls[active_choice].extend(next_p["xml_strings"])
                    if next_p["has_red"]:
                        correct_answer = active_choice
                        found_red = True
                    idx += 1
                else:
                    idx += 1

        self.part1_questions.append({
            "id": len(self.part1_questions) + 1,
            "original_num": q_num,
            "question": clean_q_text,
            "choices": choices,
            "choice_xmls": choice_xmls,
            "correct": correct_answer,
            "has_red": found_red,
            "xml_strings": clean_q_xmls
        })
        return idx

    def _extract_choices_from_p(self, p_data) -> dict:
        """Extract choices A, B, C, D from paragraph runs/text and determine which one is red."""
        res = {}
        text = p_data["raw_text"]
        fmt = p_data["formatted_text"]
        
        matches = list(re.finditer(r"(?:^|\s+)([A-D])[\.\:\)]\s*", text))
        if not matches:
            m_single = re.match(r"^\s*([A-D])[\.\:\)]\s*(.*)", text)
            if m_single:
                key = m_single.group(1).upper()
                c_text = re.sub(r"^\s*([A-D])[\.\:\)]\s*", "", fmt).strip()
                c_xmls = strip_elements_prefix(p_data["xml_strings"], r"^\s*[A-D][\.\:\)]\s*")
                res[key] = {
                    "text": c_text,
                    "is_red": p_data["has_red"],
                    "xml_strings": c_xmls
                }
            return res
            
        # Build character-level or run-level color mapping
        # Each element in oxml_elements has a text segment and is_red flag
        elem_spans = []
        curr_pos = 0
        for elem in p_data["oxml_elements"]:
            txt, is_red = extract_element_text_with_formatting(elem)
            clean_t = re.sub(r"<[^>]+>", "", txt)
            start = curr_pos
            end = curr_pos + len(clean_t)
            elem_spans.append((start, end, is_red, elem, txt))
            curr_pos = end

        for i, match in enumerate(matches):
            key = match.group(1).upper()
            m_start = match.start(1)
            m_end = matches[i+1].start(1) if i + 1 < len(matches) else len(text)
            
            # Substring text: from after marker to m_end
            c_text_raw = re.sub(r"^[A-D][\.\:\)]\s*", "", text[m_start:m_end]).strip()
            
            # Check if any run overlapping [m_start, m_end] is red
            choice_is_red = False
            choice_xmls = []
            for s, e, is_red, elem, fmt_txt in elem_spans:
                # Check overlap between [s, e] and [m_start, m_end]
                if max(s, m_start) < min(e, m_end):
                    if is_red:
                        choice_is_red = True
                    choice_xmls.append(elem)
                    
            c_clean_xmls = strip_elements_prefix(serialize_oxml_elements(choice_xmls), r"^\s*[A-D][\.\:\)]\s*")
            res[key] = {
                "text": c_text_raw,
                "is_red": choice_is_red,
                "xml_strings": c_clean_xmls
            }
        return res

    # ================= PART 2 PARSING (TRUE / FALSE) =================
    def _parse_part2_question(self, paragraphs_data, start_idx) -> int:
        p_data = paragraphs_data[start_idx]
        raw_text = p_data["raw_text"]
        q_match = QUESTION_REGEX.search(raw_text)
        q_num = int(q_match.group(2)) if q_match else len(self.part2_questions) + 1
        
        clean_q_text = re.sub(r"^\s*(Câu|Bài)\s+\d+[\.\:\-\s]+", "", p_data["formatted_text"]).strip()
        clean_q_xmls = strip_elements_prefix(p_data["xml_strings"], r"^\s*(Câu|Bài)\s+\d+[\.\:\-\s]*")
        items = {
            "a": {"text": "", "correct": False},
            "b": {"text": "", "correct": False},
            "c": {"text": "", "correct": False},
            "d": {"text": "", "correct": False}
        }
        last_key = None
        idx = start_idx + 1
        while idx < len(paragraphs_data):
            next_p = paragraphs_data[idx]
            next_text = next_p["raw_text"].strip()
            
            if (QUESTION_REGEX.search(next_text) or PART1_REGEX.search(next_text) or 
                PART2_REGEX.search(next_text) or PART3_REGEX.search(next_text) or 
                PART4_REGEX.search(next_text) or 
                re.search(r"(BẢNG\s+ĐÁP\s+ÁN|HƯỚNG\s+DẪN\s+CHẤM|\-\-\-+\s*HẾT)", next_text, re.IGNORECASE)):
                break
                
            # Check if matches a), b), c), d) or a., b., c., d.
            item_match = re.match(r"^\s*([a-d])[\)\.]\s*(.*)", next_text, re.IGNORECASE)
            if item_match:
                key = item_match.group(1).lower()
                last_key = key
                content = re.sub(r"^\s*([a-d])[\)\.]\s*", "", next_p["formatted_text"]).strip()
                
                # Check True or False:
                # 1. Has red color? In Vietnam tests, true items or [Đúng] are colored red
                is_true = next_p["has_red"]
                if "[ĐÚNG]" in next_text.upper() or "(ĐÚNG)" in next_text.upper():
                    is_true = True
                elif "[SAI]" in next_text.upper() or "(SAI)" in next_text.upper():
                    is_true = False
                    
                # Clean out [Đúng] or [Sai] tags from the visible text and XML
                clean_content = re.sub(r"\[(Đúng|Sai)\]|\((Đúng|Sai)\)", "", content, flags=re.IGNORECASE).strip()
                item_clean_xmls = strip_elements_prefix(next_p["xml_strings"], r"^\s*[a-d][\)\.]\s*")
                item_clean_xmls, _ = strip_elements_suffix(item_clean_xmls, r"(\[(Đúng|Sai)\]|\((Đúng|Sai)\))")
                
                items[key] = {
                    "text": clean_content,
                    "correct": is_true, # True = Đúng, False = Sai
                    "xml_strings": item_clean_xmls
                }
                idx += 1
            else:
                if last_key is None:
                    # Additional paragraph of the stem before any a), b), c), d)
                    clean_q_text += "<br>" + next_p["formatted_text"]
                    clean_q_xmls.extend(next_p["xml_strings"])
                    idx += 1
                else:
                    # Multi-paragraph/continuation of the current item (e.g. item a)
                    is_true_cont = next_p["has_red"]
                    if "[ĐÚNG]" in next_text.upper() or "(ĐÚNG)" in next_text.upper():
                        items[last_key]["correct"] = True
                    elif "[SAI]" in next_text.upper() or "(SAI)" in next_text.upper():
                        items[last_key]["correct"] = False
                    elif is_true_cont:
                        items[last_key]["correct"] = True

                    clean_cont = re.sub(r"\[(Đúng|Sai)\]|\((Đúng|Sai)\)", "", next_p["formatted_text"], flags=re.IGNORECASE).strip()
                    cont_xmls, _ = strip_elements_suffix(next_p["xml_strings"], r"(\[(Đúng|Sai)\]|\((Đúng|Sai)\))")
                    if clean_cont:
                        items[last_key]["text"] = (items[last_key]["text"] + " " + clean_cont).strip()
                        items[last_key]["xml_strings"].extend(cont_xmls)
                    idx += 1

        self.part2_questions.append({
            "id": len(self.part2_questions) + 1,
            "original_num": q_num,
            "question": clean_q_text,
            "items": items,
            "xml_strings": clean_q_xmls
        })
        return idx

    # ================= PART 3 PARSING (SHORT ANSWER) =================
    def _parse_part3_question(self, paragraphs_data, start_idx) -> int:
        p_data = paragraphs_data[start_idx]
        raw_text = p_data["raw_text"]
        q_match = QUESTION_REGEX.search(raw_text)
        q_num = int(q_match.group(2)) if q_match else len(self.part3_questions) + 1
        
        full_text = p_data["formatted_text"]
        full_raw = raw_text
        clean_q_xmls = strip_elements_prefix(p_data["xml_strings"], r"^\s*(Câu|Bài)\s+\d+[\.\:\-\s]*")
        
        idx = start_idx + 1
        while idx < len(paragraphs_data):
            next_p = paragraphs_data[idx]
            next_text = next_p["raw_text"].strip()
            if (QUESTION_REGEX.search(next_text) or PART1_REGEX.search(next_text) or 
                PART2_REGEX.search(next_text) or PART3_REGEX.search(next_text) or 
                PART4_REGEX.search(next_text) or 
                re.search(r"(BẢNG\s+ĐÁP\s+ÁN|HƯỚNG\s+DẪN\s+CHẤM|\-\-\-+\s*HẾT)", next_text, re.IGNORECASE)):
                break
            full_text += "<br>" + next_p["formatted_text"]
            clean_q_xmls.extend(next_p["xml_strings"])
            full_raw += "\n" + next_text
            idx += 1

        # Extract answer: either after "Đáp án:" or red runs
        clean_q = full_text
        answer = ""
        ans_pattern = (
            r"(?:"
            r"(?:<br>|\n|^|\s{2,}|\b)(?:Đáp\s*án)\s*[:=]?\s*([^\n<]+)|"
            r"(?:<br>|\n|^|\s{2,}|\b)(?:Đ\/[aA]|ĐA|Trả\s*lời)\s*[:=]\s*([^\n<]+)|"
            r"(?:<br>|\n|^)\s*(?:Kết\s*quả|KQ)\s*[:=]\s*([^\n<?]+)"
            r")"
        )
        ans_match = re.search(ans_pattern, full_text, re.IGNORECASE)
        if ans_match:
            cand_ans = (ans_match.group(1) or ans_match.group(2) or ans_match.group(3) or "").strip()
            if cand_ans and not any(k in cand_ans.upper() for k in ["CÂU ", "PHẦN ", "BÀI "]):
                answer = cand_ans
            # Remove ONLY the matched answer line from question text
            clean_q = (full_text[:ans_match.start()] + full_text[ans_match.end():]).strip()
            clean_q = re.sub(r"(<br>)+$", "", clean_q).strip()
        else:
            # Clean bare "Đáp án:" if it was an empty student box at end of question
            clean_q = re.sub(r"(?:<br>|\n)?\s*(?:Đáp\s*án\s*[:=]?|(?:Đ\/[aA]|ĐA|Trả\s*lời)\s*[:=])\s*$", "", clean_q, flags=re.IGNORECASE).strip()
            # Check red elements in paragraph
            if p_data["has_red"]:
                for run_elem in p_data["oxml_elements"]:
                    txt, is_red = extract_element_text_with_formatting(run_elem)
                    if is_red and txt.strip():
                        answer = txt.strip()
                        break

        # Strip answer from clean_q_xmls so student exam NEVER contains the answer
        safe_suffix_pattern = r"((?:^|\n|\s{2,}|\b)(?:Đáp\s*án\s*[:=]?|(?:Đ\/[aA]|ĐA|Trả\s*lời)\s*[:=]))"
        clean_q_xmls, removed_ans = strip_elements_suffix(clean_q_xmls, safe_suffix_pattern)
        if removed_ans and not answer:
            m_a = re.search(r"(?:Đáp\s*án|Đ\/[aA]|ĐA|Trả\s*lời)\s*[:=]?\s*([^\n<]+)", removed_ans, re.IGNORECASE)
            if m_a:
                answer = m_a.group(1).strip()

        # Remove "Câu X:"
        clean_q = re.sub(r"^\s*(Câu|Bài)\s+\d+[\.\:\-\s]+", "", clean_q).strip()

        self.part3_questions.append({
            "id": len(self.part3_questions) + 1,
            "original_num": q_num,
            "question": clean_q,
            "answer": answer,
            "xml_strings": clean_q_xmls
        })
        return idx

    # ================= PART 4 PARSING (ESSAY) =================
    def _parse_part4_question(self, paragraphs_data, start_idx) -> int:
        p_data = paragraphs_data[start_idx]
        raw_text = p_data["raw_text"]
        q_match = QUESTION_REGEX.search(raw_text)
        q_num = int(q_match.group(2)) if q_match else len(self.part4_questions) + 1
        
        full_text = p_data["formatted_text"]
        clean_q_xmls = strip_elements_prefix(p_data["xml_strings"], r"^\s*(Câu|Bài)\s+\d+(\s*\([^\)]+\))?[\.\:\-\s]*")
        
        idx = start_idx + 1
        while idx < len(paragraphs_data):
            next_p = paragraphs_data[idx]
            next_text = next_p["raw_text"].strip()
            if (QUESTION_REGEX.search(next_text) or PART1_REGEX.search(next_text) or 
                PART2_REGEX.search(next_text) or PART3_REGEX.search(next_text) or 
                PART4_REGEX.search(next_text) or 
                re.search(r"(BẢNG\s+ĐÁP\s+ÁN|HƯỚNG\s+DẪN\s+CHẤM|\-\-\-+\s*HẾT)", next_text, re.IGNORECASE)):
                break
            full_text += "<br>" + next_p["formatted_text"]
            clean_q_xmls.extend(next_p["xml_strings"])
            idx += 1

        # Separate question and grading guide / answer if marked
        parts = re.split(r"((?:<br>|\n|^)\s*(?:Hướng\s*dẫn\s*chấm|HD\s*chấm|Lời\s*giải|HƯỚNG\s*DẪN\s*CHẤM|Đáp\s*án)\s*[:=\n])", full_text, flags=re.IGNORECASE)
        if len(parts) >= 3:
            q_part = parts[0].strip()
            guide_part = "".join(parts[1:]).strip()
        else:
            q_part = full_text
            guide_part = ""

        # Strip grading guide from clean_q_xmls so student exam NEVER contains solutions
        clean_q_xmls, removed_guide = strip_elements_suffix(clean_q_xmls, r"((?:^|\n|\s{2,}|\b)(?:Hướng\s*dẫn\s*chấm|HD\s*chấm|Lời\s*giải|HƯỚNG\s*DẪN\s*CHẤM|Đáp\s*án)\s*[\:\=\s\n])")
        if removed_guide and not guide_part:
            guide_part = removed_guide.strip()

        # Remove "Câu X (X điểm):"
        clean_q = re.sub(r"^\s*(Câu|Bài)\s+\d+(\s*\([^\)]+\))?[\.\:\-\s]+", "", q_part).strip()
        clean_q = re.sub(r"(\<br\>)?\s*(?:Hướng\s*dẫn\s*chấm|HD\s*chấm|Lời\s*giải)\s*[\:\=\s\n]+.*", "", clean_q, flags=re.IGNORECASE | re.DOTALL).strip()

        self.part4_questions.append({
            "id": len(self.part4_questions) + 1,
            "original_num": q_num,
            "question": clean_q,
            "guide": guide_part,
            "xml_strings": clean_q_xmls
        })
        return idx

    # ================= TRAILING ANSWER KEY PARSER =================
    def _parse_trailing_answers(self, paragraphs_data, start_idx: int):
        """
        Parses answers and grading rubrics from the separate BẢNG ĐÁP ÁN & HƯỚNG DẪN CHẤM
        at the end of the document.
        """
        self._check_trailing_tables()

        # Also scan text paragraphs from start_idx onwards for text answers and Part 4 rubric
        curr_p4_qnum = None
        for p in paragraphs_data[start_idx:]:
            txt = p["raw_text"].strip()
            fmt = p["formatted_text"].strip()
            if not txt:
                continue

            # Check text matches for Part 1: "Câu 1: B" or "1. B" or "1 - B"
            p1_matches = re.findall(r"(?:^|\s|Câu\s*)(\d+)[\.\:\s\-]+([A-D])(?:\s|$|[\,\;])", txt, re.IGNORECASE)
            for m_n, m_a in p1_matches:
                q_n = int(m_n)
                if q_n <= len(self.part1_questions):
                    self._set_part1_answer(q_n, m_a.upper())

            # Check text matches for Part 2: "Câu 1: a) Đúng | b) Đúng | c) Sai | d) Sai"
            p2_matches = re.findall(r"(?:^|\s|Ý\s*)([a-d])[\)\.\:\s\-]+(Đ|Đúng|True|S|Sai|False)\b", txt, re.IGNORECASE)
            m_q2 = re.search(r"Câu\s*(\d+)", txt, re.IGNORECASE)
            if p2_matches and m_q2:
                q_n = int(m_q2.group(1))
                for itm_k, val_str in p2_matches:
                    is_tr = (val_str.upper() in ["Đ", "ĐÚNG", "TRUE"])
                    self._set_part2_answer(q_n, itm_k.lower(), is_tr)

            # Check Part 4 question guides: "Câu 1 (1,5 điểm): ..."
            m_p4 = re.search(r"^\s*Câu\s*(\d+)\s*(?:\([^\)]+\))?[\:\.\-]\s*(.*)", txt, re.IGNORECASE)
            if m_p4 and len(self.part4_questions) > 0:
                curr_p4_qnum = int(m_p4.group(1))
                rest = m_p4.group(2).strip()
                if rest:
                    self._append_part4_guide(curr_p4_qnum, rest)
            elif curr_p4_qnum is not None and any(sym in txt for sym in ["•", "-", "Diện tích", "Chiều cao", "Thể tích", "Vận tốc", "Quãng đường"]):
                if len(self.part4_questions) > 0:
                    self._append_part4_guide(curr_p4_qnum, fmt)

    def _check_trailing_tables(self):
        """Scans tables in docx.Document.tables to extract Part 1, Part 2, and Part 3 answers."""
        for tbl in self.doc.tables:
            if len(tbl.rows) < 2:
                continue
            r0_texts = [re.sub(r"\s+", " ", c.text).strip().upper() for c in tbl.rows[0].cells]
            
            # --- Check Part 1 Answer Table (Horizontal) ---
            # Row 0: "Câu", "1", "2", "3"... Row 1: "Chọn", "B", "B", "B"...
            if any(t in ["CÂU", "CÂU HỎI"] for t in r0_texts):
                nums_in_row0 = [int(re.search(r"\d+", t).group(0)) for t in r0_texts if re.search(r"^\d+$", t)]
                if len(nums_in_row0) >= 2 and len(tbl.rows) >= 2:
                    for col_idx, cell in enumerate(tbl.rows[0].cells):
                        m_num = re.search(r"^(\d+)$", cell.text.strip())
                        if m_num and col_idx < len(tbl.rows[1].cells):
                            q_num = int(m_num.group(1))
                            ans_val = tbl.rows[1].cells[col_idx].text.strip().upper()
                            if ans_val in ["A", "B", "C", "D"]:
                                self._set_part1_answer(q_num, ans_val)
                    continue

                # --- Check Part 2 Table ---
                # Columns: ["Câu", "Lệnh hỏi", "Đáp án"]
                if any("LỆNH HỎI" in t or "Ý" in t for t in r0_texts) and any("ĐÁP ÁN" in t for t in r0_texts):
                    for row in tbl.rows[1:]:
                        r_txts = [c.text.strip() for c in row.cells]
                        if len(r_txts) >= 3:
                            m_q = re.search(r"(\d+)", r_txts[0])
                            m_itm = re.search(r"([a-d])", r_txts[1], re.IGNORECASE)
                            ans_str = r_txts[2].strip().upper()
                            if m_q and m_itm:
                                q_n = int(m_q.group(1))
                                itm_k = m_itm.group(1).lower()
                                is_true = (ans_str in ["Đ", "ĐÚNG", "TRUE", "T", "1"])
                                self._set_part2_answer(q_n, itm_k, is_true)
                    continue

                # --- Check Part 3 Table (Vertical: Câu | Đáp án) ---
                if any("ĐÁP ÁN" in t for t in r0_texts) and len(tbl.columns) == 2:
                    for row in tbl.rows[1:]:
                        q_str = row.cells[0].text.strip()
                        ans_str = row.cells[1].text.strip()
                        m_q = re.search(r"(\d+)", q_str)
                        if m_q:
                            q_n = int(m_q.group(1))
                            if ans_str.upper() in ["A", "B", "C", "D"] and q_n <= len(self.part1_questions) and not any(q.get("has_red") for q in self.part1_questions):
                                self._set_part1_answer(q_n, ans_str.upper())
                            else:
                                self._set_part3_answer(q_n, ans_str)
                    continue

    def _set_part1_answer(self, q_num: int, ans: str):
        for q in self.part1_questions:
            if q.get("original_num") == q_num or q.get("id") == q_num:
                q["correct"] = ans
                q["has_red"] = True
                break

    def _set_part2_answer(self, q_num: int, item_key: str, is_true: bool):
        for q in self.part2_questions:
            if q.get("original_num") == q_num or q.get("id") == q_num:
                if item_key in q.get("items", {}):
                    q["items"][item_key]["correct"] = is_true
                break

    def _set_part3_answer(self, q_num: int, ans: str):
        for q in self.part3_questions:
            if q.get("original_num") == q_num or q.get("id") == q_num:
                q["answer"] = ans
                break

    def _append_part4_guide(self, q_num: int, text: str):
        for q in self.part4_questions:
            if q.get("original_num") == q_num or q.get("id") == q_num:
                if q.get("guide"):
                    if text not in q["guide"]:
                        q["guide"] += "<br>" + text
                else:
                    q["guide"] = text
                break
