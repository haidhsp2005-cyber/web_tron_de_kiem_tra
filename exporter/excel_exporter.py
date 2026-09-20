import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelExporter:
    def __init__(self, shuffled_result: dict, output_path: str = "exports/Bang_dap_an_tong_hop.xlsx"):
        self.result = shuffled_result
        self.metadata = shuffled_result.get("metadata", {})
        self.summary = shuffled_result.get("summary", {})
        self.output_path = output_path
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

    def export(self) -> str:
        wb = openpyxl.Workbook()
        # Remove default sheet
        default_sheet = wb.active
        
        codes = self.summary.get("codes", [])
        
        # Style helpers
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )
        
        font_header = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Times New Roman", size=11)
        font_data_bold = Font(name="Times New Roman", size=11, bold=True)
        
        align_center = Alignment(horizontal="center", vertical="center")
        
        # 1. SHEET PHẦN I
        ws1 = wb.create_sheet(title="Phần I - Trắc nghiệm")
        ws1.views.sheetView[0].showGridLines = True
        
        # Title
        ws1.merge_cells("A1:E1")
        ws1["A1"] = f"BẢNG ĐÁP ÁN PHẦN I - TRẮC NGHIỆM KHÁCH QUAN (MÔN: {self.metadata.get('subject', 'KHTN').upper()})"
        ws1["A1"].font = Font(name="Times New Roman", size=13, bold=True, color="1E40AF")
        ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")
        
        # Headers
        ws1.cell(row=3, column=1, value="Câu").font = font_header
        ws1.cell(row=3, column=1).fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        ws1.cell(row=3, column=1).alignment = align_center
        ws1.cell(row=3, column=1).border = thin_border
        
        for col_idx, c in enumerate(codes, start=2):
            cell = ws1.cell(row=3, column=col_idx, value=f"Mã {c}")
            cell.font = font_header
            cell.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
            cell.alignment = align_center
            cell.border = thin_border
            
        p1_rows = self.summary.get("part1", {}).get("rows", [])
        for row_idx, r_data in enumerate(p1_rows, start=4):
            bg_color = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
            fill_row = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            
            c_num = ws1.cell(row=row_idx, column=1, value=f"Câu {r_data['question_num']}")
            c_num.font = font_data_bold
            c_num.fill = fill_row
            c_num.alignment = align_center
            c_num.border = thin_border
            
            for col_idx, c in enumerate(codes, start=2):
                ans = r_data.get(c, "")
                cell = ws1.cell(row=row_idx, column=col_idx, value=ans)
                cell.font = font_data_bold
                cell.fill = fill_row
                cell.alignment = align_center
                cell.border = thin_border

        self._auto_column_width(ws1)

        # 2. SHEET PHẦN II
        ws2 = wb.create_sheet(title="Phần II - Đúng Sai")
        ws2.views.sheetView[0].showGridLines = True
        
        ws2.merge_cells("A1:E1")
        ws2["A1"] = f"BẢNG ĐÁP ÁN PHẦN II - CÂU HỎI ĐÚNG / SAI (MÔN: {self.metadata.get('subject', 'KHTN').upper()})"
        ws2["A1"].font = Font(name="Times New Roman", size=13, bold=True, color="047857")
        ws2["A1"].alignment = Alignment(horizontal="left", vertical="center")
        
        ws2.cell(row=3, column=1, value="Lệnh hỏi").font = font_header
        ws2.cell(row=3, column=1).fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
        ws2.cell(row=3, column=1).alignment = align_center
        ws2.cell(row=3, column=1).border = thin_border
        
        for col_idx, c in enumerate(codes, start=2):
            cell = ws2.cell(row=3, column=col_idx, value=f"Mã {c}")
            cell.font = font_header
            cell.fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
            cell.alignment = align_center
            cell.border = thin_border
            
        p2_rows = self.summary.get("part2", {}).get("rows", [])
        for row_idx, r_data in enumerate(p2_rows, start=4):
            bg_color = "F0FDF4" if row_idx % 2 == 0 else "FFFFFF"
            fill_row = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            
            c_lbl = ws2.cell(row=row_idx, column=1, value=r_data["label"])
            c_lbl.font = font_data_bold
            c_lbl.fill = fill_row
            c_lbl.alignment = align_center
            c_lbl.border = thin_border
            
            for col_idx, c in enumerate(codes, start=2):
                ans = r_data.get(c, "")
                cell = ws2.cell(row=row_idx, column=col_idx, value=ans)
                cell.font = Font(name="Times New Roman", size=11, bold=True, color="047857" if ans == "Đ" else "4B5563")
                cell.fill = fill_row
                cell.alignment = align_center
                cell.border = thin_border

        self._auto_column_width(ws2)

        # 3. SHEET PHẦN III
        ws3 = wb.create_sheet(title="Phần III - Trả lời ngắn")
        ws3.views.sheetView[0].showGridLines = True
        
        ws3.merge_cells("A1:E1")
        ws3["A1"] = f"BẢNG ĐÁP ÁN PHẦN III - TRẢ LỜI NGẮN (MÔN: {self.metadata.get('subject', 'KHTN').upper()})"
        ws3["A1"].font = Font(name="Times New Roman", size=13, bold=True, color="B45309")
        ws3["A1"].alignment = Alignment(horizontal="left", vertical="center")
        
        ws3.cell(row=3, column=1, value="Câu").font = font_header
        ws3.cell(row=3, column=1).fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
        ws3.cell(row=3, column=1).alignment = align_center
        ws3.cell(row=3, column=1).border = thin_border
        
        for col_idx, c in enumerate(codes, start=2):
            cell = ws3.cell(row=3, column=col_idx, value=f"Mã {c}")
            cell.font = font_header
            cell.fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
            cell.alignment = align_center
            cell.border = thin_border
            
        p3_rows = self.summary.get("part3", {}).get("rows", [])
        for row_idx, r_data in enumerate(p3_rows, start=4):
            bg_color = "FFFBEB" if row_idx % 2 == 0 else "FFFFFF"
            fill_row = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            
            c_num = ws3.cell(row=row_idx, column=1, value=f"Câu {r_data['question_num']}")
            c_num.font = font_data_bold
            c_num.fill = fill_row
            c_num.alignment = align_center
            c_num.border = thin_border
            
            for col_idx, c in enumerate(codes, start=2):
                ans = str(r_data.get(c, ""))
                cell = ws3.cell(row=row_idx, column=col_idx, value=ans)
                cell.font = font_data_bold
                cell.fill = fill_row
                cell.alignment = align_center
                cell.border = thin_border

        self._auto_column_width(ws3)

        # Remove default sheet if exists
        if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
            wb.remove(default_sheet)

        wb.save(self.output_path)
        return self.output_path

    def _auto_column_width(self, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len and not cell.coordinate.startswith("A1"):
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 5, 12)
